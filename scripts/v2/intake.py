"""Archive quote originals and extract source evidence, never publish inferred prices.

XLSX cells are streamed from XML, so formatted blank rows do not truncate prices.
Numbers, cached values and formulas remain distinct; numeric date serials stay raw.
"""
from __future__ import annotations
import argparse
import gzip
import hashlib
import json
import posixpath
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

LEGACY_DEPS = Path(__file__).resolve().parents[2] / 'analysis' / '.deps'
if LEGACY_DEPS.exists():
    sys.path.insert(0, str(LEGACY_DEPS))

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def write_record(stream, record):
    stream.write(json.dumps(record, ensure_ascii=False, default=str) + '\n')


def xlsx_evidence(path, stream):
    sheets = []
    with zipfile.ZipFile(path) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            with archive.open('xl/sharedStrings.xml') as source:
                for _, element in ET.iterparse(source, events=('end',)):
                    if element.tag == NS + 'si':
                        strings.append(''.join(node.text or '' for node in element.iter(NS + 't')))
                        element.clear()
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {r.attrib['Id']: r.attrib['Target'] for r in relations}
        for sheet in workbook.find(NS + 'sheets'):
            name = sheet.attrib['name']
            target = targets[sheet.attrib[REL]]
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            count, merges = 0, []
            with archive.open(target) as source:
                for _, element in ET.iterparse(source, events=('end',)):
                    if element.tag == NS + 'mergeCell':
                        merges.append(element.attrib.get('ref'))
                    if element.tag != NS + 'c':
                        if element.tag == NS + 'row':
                            element.clear()
                        continue
                    value_node, formula_node = element.find(NS + 'v'), element.find(NS + 'f')
                    kind = element.attrib.get('t', 'n')
                    raw = value_node.text if value_node is not None else None
                    if kind == 's' and raw is not None:
                        value = strings[int(raw)]
                    elif kind == 'inlineStr':
                        value = ''.join(n.text or '' for n in element.iter(NS + 't'))
                    else:
                        value = raw
                    if value not in (None, '') or formula_node is not None:
                        item = {'sheet': name, 'cell': element.attrib['r'], 'type': kind,
                                'value': value, 'style': element.attrib.get('s')}
                        if formula_node is not None:
                            item['formula'] = formula_node.text
                            item['formula_attributes'] = formula_node.attrib
                            item['cached_value'] = raw
                        write_record(stream, item)
                        count += 1
                    element.clear()
            sheets.append({'name': name, 'state': sheet.attrib.get('state', 'visible'),
                           'cells': count, 'merged_ranges': merges})
    return sheets


def xls_evidence(path, stream):
    import xlrd
    from openpyxl.utils import get_column_letter
    book = xlrd.open_workbook(str(path), on_demand=True, formatting_info=True)
    sheets = []
    try:
        for index, sheet in enumerate(book.sheets()):
            count = 0
            for row in range(sheet.nrows):
                for column in range(sheet.ncols):
                    cell = sheet.cell(row, column)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    write_record(stream, {'sheet': sheet.name,
                        'cell': f'{get_column_letter(column + 1)}{row + 1}',
                        'type': cell.ctype, 'value': cell.value})
                    count += 1
            sheets.append({'name': sheet.name, 'cells': count,
                'state': 'hidden' if getattr(book, '_sheet_visibility', [0] * book.nsheets)[index] else 'visible',
                'merged_ranges': sheet.merged_cells,
                'limitation': 'BIFF formula text unavailable; values are stored results'})
    finally:
        book.release_resources()
    return sheets


def extract(path, output):
    if path.suffix.lower() == '.pdf':
        from pypdf import PdfReader
        pages = []
        with gzip.open(output, 'wt', encoding='utf-8') as stream:
            for number, page in enumerate(PdfReader(path).pages, 1):
                text = page.extract_text() or ''
                write_record(stream, {'page': number, 'text': text})
                pages.append({'page': number, 'characters': len(text)})
        return {'pages': pages, 'limitation': 'Text evidence only; table layout and prices require review'}
    with gzip.open(output, 'wt', encoding='utf-8') as stream:
        sheets = xls_evidence(path, stream) if path.suffix.lower() == '.xls' else xlsx_evidence(path, stream)
    return {'sheets': sheets, 'cells': sum(sheet['cells'] for sheet in sheets)}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path('.local/v2/intake'))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    manifest = {'version': 1, 'cost_basis_confirmed': True, 'source_root': str(args.root.resolve()),
                'generated_at': datetime.now(timezone.utc).isoformat(), 'files': []}
    for original in sorted(args.root.iterdir()):
        if original.suffix.lower() not in ('.xlsx', '.xls', '.xlsm', '.pdf') or original.name.startswith('~$'):
            continue
        sha = digest(original)
        folder = args.output / sha
        folder.mkdir(exist_ok=True)
        stored = folder / ('original' + original.suffix.lower())
        if not stored.exists():
            shutil.copy2(original, stored)
        if digest(stored) != sha:
            raise ValueError(f'Archive checksum mismatch: {original.name}')
        record = {'id': sha, 'sha256': sha, 'name': original.name, 'size_bytes': stored.stat().st_size,
                  'original_path': str(original.resolve()), 'storage_path': str(stored.resolve()),
                  'evidence_path': str((folder / 'evidence.jsonl.gz').resolve()),
                  'pricing_readiness': 'needs_review', 'cost_basis': 'procurement', 'status': 'extracted'}
        metadata = folder / 'extraction.json'
        try:
            if metadata.exists():
                evidence = json.loads(metadata.read_text(encoding='utf-8'))
            else:
                evidence = extract(stored, folder / 'evidence.jsonl.gz')
                metadata.write_text(json.dumps(evidence, ensure_ascii=False), encoding='utf-8')
            record.update(evidence)
        except Exception as error:
            record.update(status='failed', error=f'{type(error).__name__}: {error}')
        manifest['files'].append(record)
        print(original.name, record['status'], record.get('cells', ''), flush=True)
        # Preserve completed work if a later source fails or the process is interrupted.
        (args.output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'files': len(manifest['files']), 'failed': sum(f['status'] == 'failed' for f in manifest['files'])}))


if __name__ == '__main__':
    main()
