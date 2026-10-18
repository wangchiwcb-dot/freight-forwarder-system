"""Independently compare every priced database tier against the actual Excel file.
This does not use the JSON extraction/normalizer as its numeric evidence.
"""
import argparse
import hashlib
import json
import sqlite3
import sys
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'analysis' / '.deps'))
import xlrd


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--database', default='.local/d1-import/freight-quotes.sqlite')
    ap.add_argument('--output', default='.local/d1-import/workbook-verification.json')
    args = ap.parse_args()
    db = sqlite3.connect(args.database)
    query = '''SELECT f.sha256,f.local_path,r.sheet_name,t.source_cell,t.price_decimal,t.id
      FROM quote_tiers t JOIN quote_versions v ON v.id=t.quote_version_id
      JOIN import_rows r ON r.quote_version_id=v.id JOIN import_batches b ON b.id=v.batch_id
      JOIN source_files f ON f.id=b.source_file_id WHERE t.price_status='priced' '''
    grouped = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    for sha, file, sheet, cell, price, tier_id in db.execute(query):
        grouped[(sha, file)][sheet][cell].append((price, tier_id))
    errors, files, checked, unique_cells = [], [], 0, 0

    def compare(raw, expected, label):
        nonlocal checked, unique_cells
        unique_cells += 1
        for price, tier_id in expected:
            checked += 1
            try:
                matches = raw is not None and Decimal(str(raw)) == Decimal(price)
            except (InvalidOperation, ValueError):
                matches = False
            if not matches:
                errors.append({'cell': label, 'tier': tier_id, 'database': price, 'workbook': str(raw)})

    for (expected_hash, filename), sheets in grouped.items():
        file = Path(filename) if filename else None
        if not file or not file.exists():
            errors.append({'file': filename, 'error': 'Original workbook unavailable'})
            continue
        actual_hash = hashlib.sha256(file.read_bytes()).hexdigest()
        if expected_hash != actual_hash:
            errors.append({'file': filename, 'error': 'Source hash changed'})
            continue
        before = checked
        if file.suffix.lower() == '.xls':
            workbook = xlrd.open_workbook(str(file), on_demand=True)
            for name, cells in sheets.items():
                sh = workbook.sheet_by_name(name)
                for coord, values in cells.items():
                    r, c = openpyxl.utils.cell.coordinate_to_tuple(coord)
                    cell = sh.cell(r - 1, c - 1)
                    value = xlrd.error_text_from_code.get(cell.value, '#ERROR!') if cell.ctype == xlrd.XL_CELL_ERROR else cell.value
                    compare(value, values, f'{file.name}/{name}!{coord}')
            workbook.release_resources()
        else:
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=True, keep_links=False)
            for name, cells in sheets.items():
                ws = workbook[name]
                coordinates = {openpyxl.utils.cell.coordinate_to_tuple(coord): (coord, values) for coord, values in cells.items()}
                max_row = max(r for r, c in coordinates)
                max_col = max(c for r, c in coordinates)
                for r, row in enumerate(ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True), 1):
                    for c, raw in enumerate(row, 1):
                        if (r, c) in coordinates:
                            coord, values = coordinates[(r, c)]
                            compare(raw, values, f'{file.name}/{name}!{coord}')
            workbook.close()
        files.append({'file': str(file), 'sha256': expected_hash, 'checkedPrices': checked - before})
        print(f'{file.name}: {checked - before} prices checked', flush=True)
    expected_count = db.execute("SELECT count(*) FROM quote_tiers WHERE price_status='priced'").fetchone()[0]
    if checked != expected_count:
        errors.append({'error': 'Not all priced tiers checked', 'checked': checked, 'expected': expected_count})
    report = {'passed': not errors, 'checkedPrices': checked, 'uniqueSourceCells': unique_cells,
              'workbooks': len(files), 'files': files, 'errors': errors}
    Path(args.output).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: report[k] for k in ['passed', 'checkedPrices', 'uniqueSourceCells', 'workbooks']}, ensure_ascii=False))
    db.close()
    if errors:
        print(json.dumps(errors[:10], ensure_ascii=False))
        raise SystemExit(1)


if __name__ == '__main__':
    main()
